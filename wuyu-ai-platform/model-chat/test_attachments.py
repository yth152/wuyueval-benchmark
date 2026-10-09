import asyncio
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from fastapi.testclient import TestClient
import attachments as a
import search_tools as st
import server

def fixtures():
    from docx import Document
    from pptx import Presentation
    from openpyxl import Workbook
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    result={}
    out=io.BytesIO();doc=Document();doc.add_paragraph('验收项目：清源；现场核查 7 处。');doc.save(out);result['check.docx']=out.getvalue()
    out=io.BytesIO();ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text='清源项目';slide.placeholders[1].text='核查点位 7 处';ppt.save(out);result['check.pptx']=out.getvalue()
    out=io.BytesIO();wb=Workbook();ws=wb.active;ws.title='固废台账';ws.append(['类型','数量']);ws.append(['废包装',7]);wb.save(out);result['check.xlsx']=out.getvalue()
    out=io.BytesIO();writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 250 Td (Inspection count: 7) Tj ET');page[NameObject('/Contents')]=writer._add_object(stream);writer.write(out);result['check.pdf']=out.getvalue()
    out=io.BytesIO();Image.new('RGB',(120,80),'blue').save(out,format='PNG');result['blue.png']=out.getvalue()
    result['check.txt']='验收编号 QY-7，核查点位 7 处。'.encode('utf-8')
    return result

class AttachmentTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.patch=patch.object(a,'UPLOAD_DIR',Path(self.temp.name));self.patch.start()
    def tearDown(self):self.patch.stop();self.temp.cleanup()

    async def test_real_office_pdf_text_and_image_parsers(self):
        for name,data in fixtures().items():
            with self.subTest(name=name):
                item=await a.save_file(name,data);stored=a.read_file(item['id'])
                self.assertNotIn('text',item);self.assertNotIn('image',item)
                if name.endswith('.png'):
                    self.assertEqual(item['kind'],'image');self.assertTrue((a.UPLOAD_DIR/(item['id']+'.jpg')).exists())
                else:self.assertIn('7',stored['text'])

    async def test_history_files_stay_on_correct_turn_and_images_use_native_parts(self):
        data=fixtures();doc=await a.save_file('check.txt',data['check.txt']);image=await a.save_file('blue.png',data['blue.png'])
        body=server.ChatRequest(messages=[{'role':'user','content':'看附件','file_ids':[doc['id'],image['id']]},{'role':'assistant','content':'收到'},{'role':'user','content':'核查点位多少？'}])
        messages,notices=a.prepare_messages(body)
        self.assertIn('QY-7',messages[0]['content'][0]['text']);self.assertTrue(messages[0]['content'][2]['image_url']['url'].startswith('data:image/jpeg;base64,'))
        self.assertEqual(messages[-1]['content'],'核查点位多少？');self.assertEqual(notices,[])

    async def test_large_document_is_explicitly_truncated(self):
        item=await a.save_file('long.txt',('信息'*16000).encode())
        body=server.ChatRequest(messages=[{'role':'user','content':'总结','file_ids':[item['id']]}])
        messages,notices=a.prepare_messages(body)
        self.assertIn('仅选取前 12,000',notices[0]);self.assertLess(len(messages[0]['content']),12500)

    async def test_pdf_without_text_layer_is_not_reported_as_parsed(self):
        from pypdf import PdfWriter
        out=io.BytesIO();writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.write(out)
        with self.assertRaisesRegex(ValueError,'没有可读文字'):await a.save_file('scan.pdf',out.getvalue())

    async def test_transparent_images_keep_a_readable_white_background(self):
        out=io.BytesIO();Image.new('RGBA',(32,32),(0,0,0,0)).save(out,format='PNG')
        item=await a.save_file('transparent.png',out.getvalue())
        with Image.open(a.UPLOAD_DIR/(item['id']+'.jpg')) as image:self.assertEqual(image.getpixel((0,0)),(255,255,255))

    async def test_image_cache_namespace_preserves_legacy_and_new_pixels_and_storage(self):
        import base64
        out=io.BytesIO();Image.new('RGB',(1280,867),'red').save(out,format='JPEG',quality=88,optimize=True)
        original=out.getvalue();file_id='b'*32
        (a.UPLOAD_DIR/(file_id+'.jpg')).write_bytes(original)
        legacy={'id':file_id,'name':'old.jpg','kind':'image','size':len(original)}
        (a.UPLOAD_DIR/(file_id+'.json')).write_text(json.dumps(legacy),encoding='utf-8')
        body=server.ChatRequest(messages=[{'role':'user','content':'看图','file_ids':[file_id]}])
        messages,_=a.prepare_messages(body)
        normalized=base64.b64decode(messages[0]['content'][2]['image_url']['url'].split(',',1)[1])
        self.assertNotEqual(normalized,original)
        self.assertEqual((a.UPLOAD_DIR/(file_id+'.jpg')).read_bytes(),original)
        with Image.open(io.BytesIO(normalized)) as decoded:
            self.assertEqual(decoded.size,(1280,867));self.assertEqual(decoded.mode,'RGB')
            with Image.open(io.BytesIO(original)) as before:
                self.assertEqual(decoded.tobytes(),before.tobytes())
        self.assertEqual(a.model_image_bytes(legacy),normalized)
        fresh=await a.save_file('new.jpg',original)
        stored=a.read_file(fresh['id'])
        self.assertEqual(stored['model_image_encoding'],'baseline-jpeg-v2')
        original=(a.UPLOAD_DIR/(fresh['id']+'.jpg')).read_bytes()
        normalized=a.model_image_bytes(stored)
        self.assertNotEqual(normalized,original)
        with Image.open(io.BytesIO(normalized)) as after, Image.open(io.BytesIO(original)) as before:
            self.assertEqual(after.tobytes(),before.tobytes())

    async def test_reject_unsupported_corrupt_and_oversized_image(self):
        for name,data in [('run.exe',b'x'),('bad.png',b'not image'),('large.png',b'x'*(a.IMAGE_MB*1048576+1))]:
            with self.assertRaises(ValueError):await a.save_file(name,data)

    async def test_multipart_api_and_preview_do_not_expose_text_or_filesystem(self):
        client=TestClient(server.app,base_url='http://127.0.0.1:8910')
        response=await asyncio.to_thread(client.post,'/api/files',files={'file':('blue.png',fixtures()['blue.png'],'image/png')})
        self.assertEqual(response.status_code,200);item=response.json()
        self.assertEqual(client.get(item['preview_url']).headers['content-type'],'image/jpeg')
        self.assertEqual(client.get('/api/files/secret/preview').status_code,404)
        self.assertEqual(client.post('/api/files',files={'file':('a.txt',b'x')},headers={'Origin':'https://evil.test'}).status_code,403)

    async def test_missing_file_and_path_traversal_fail_before_model_call(self):
        with self.assertRaises(ValueError):a.read_file('../team.dpapi')
        with self.assertRaises(ValueError):a.prepare_messages(server.ChatRequest(messages=[{'role':'user','content':'看文件','file_ids':['f'*32]}]))
        with self.assertRaises(ValueError):server.ChatRequest(messages=[{'role':'user','content':'看文件','file_ids':['a'*32]*6}])

class BilingualAndDateTests(unittest.TestCase):
    def test_requires_two_valid_language_queries(self):
        args=json.dumps({'query_zh':'固废 遥感','query_en':'solid waste remote sensing'})
        self.assertEqual(st.validate_call('scholar_search',args)['period'],'recent_10_years')
        for args in [{'query':'固废'},{'query_zh':'固废','query_en':'固废'},{'query_zh':'solid waste','query_en':'solid waste'},{'query_zh':'固废','query_en':'waste','size':100}]:
            with self.assertRaises(ValueError):st.validate_call('web_search',json.dumps(args))

    def test_recent_filter_uses_dates_and_marks_unknown(self):
        from datetime import date
        sources=[{'date':value,'title':value} for value in ['2010-01-01','2022-03-02','','2030-02-01','2016-01-01','2016-11-01','2026-12-01']]
        with patch.object(st,'search_window',return_value=(date(2016,10,5),date(2026,10,5))):kept,excluded=st.recent_sources(sources)
        self.assertEqual(excluded,4);self.assertEqual([s['date'] for s in kept],['2022-03-02','2016-11-01','']);self.assertEqual(kept[-1]['year_status'],'unknown')

if __name__=='__main__':unittest.main()
