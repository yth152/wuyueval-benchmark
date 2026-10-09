import io
import json
import re
import subprocess
import shutil
import tempfile
import zipfile
from pathlib import Path
from defusedxml.ElementTree import fromstring

TEXT_EXT = {'.txt', '.md', '.csv', '.json', '.log', '.tsv'}
SUPPORTED = TEXT_EXT | {'.docx', '.xlsx', '.pptx', '.pdf', '.doc', '.xls', '.ppt'}
TEXT_LIMIT = 180000

def safe_zip(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    entries = z.infolist()
    if len(entries) > 3000 or sum(e.file_size for e in entries) > 64 * 1024 * 1024:
        raise ValueError('文档解压后过大，请拆分后上传（最多 64 MB）。')
    if any(e.file_size > 24 * 1024 * 1024 or (e.file_size > 1024 * 1024 and e.file_size / max(e.compress_size, 1) > 200) for e in entries):
        raise ValueError('文档压缩比例异常或单个内容过大。')
    if any(e.filename.endswith('vbaProject.bin') for e in entries):
        raise ValueError('不接受包含宏的文档，请另存为无宏格式。')
    return z

def extract(name, data):
    ext = Path(name).suffix.lower()
    warning = ''
    if ext not in SUPPORTED:
        raise ValueError('支持 Word、Excel、PPT、PDF、TXT、Markdown、CSV 和 JSON。')
    if ext in TEXT_EXT:
        try: text = data.decode('utf-8-sig')
        except UnicodeDecodeError: text = data.decode('gb18030')
        if '\x00' in text: raise ValueError('文件不是可读取的文本。')
    elif ext in {'.docx', '.pptx', '.xlsx'}:
        z = safe_zip(data)
        if ext == '.docx':
            root = fromstring(z.read('word/document.xml'))
            text = '\n'.join(''.join(n.itertext()) for n in root.iter() if n.tag.endswith('}t'))
        elif ext == '.pptx':
            names = sorted([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)], key=lambda n: int(re.search(r'(\d+)\.xml', n)[1]))
            text = '\n\n'.join(f'第 {i+1} 页\n' + '\n'.join(n.text or '' for n in fromstring(z.read(p)).iter() if n.tag.endswith('}t')) for i, p in enumerate(names))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            parts, chars, count = [], 0, 0
            for ws in wb:
                parts.append('工作表：' + ws.title)
                for row in ws.iter_rows(max_row=min(ws.max_row or 10000,10000), max_col=min(ws.max_column or 100,100), values_only=True):
                    line = '\t'.join('' if v is None else str(v) for v in row)
                    if line.strip(): parts.append(line)
                    chars += len(line); count += 1
                    if chars > TEXT_LIMIT or count > 12000: break
                if chars > TEXT_LIMIT or count > 12000: break
            wb.close(); text = '\n'.join(parts)
            warning = '表格读取缓存值；公式不重新计算，图表和图片不解析。每表最多 100 列、10,000 行。'
    elif ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted: raise ValueError('请先解除 PDF 密码保护。')
        parts = []
        for i, page in enumerate(reader.pages[:100]):
            parts.append(f'第 {i+1} 页\n' + (page.extract_text() or ''))
            if sum(map(len, parts)) > TEXT_LIMIT: break
        text = '\n\n'.join(parts)
        warning = 'PDF 最多解析前 100 页的文本层；扫描图片和图表暂不做 OCR。'
    elif ext == '.xls':
        import xlrd
        wb = xlrd.open_workbook(file_contents=data, on_demand=True)
        parts=[]
        for sheet in wb.sheets():
            parts.append('工作表：' + sheet.name)
            for i in range(min(sheet.nrows, 10000)):
                parts.append('\t'.join(str(v) for v in sheet.row_values(i, end_colx=min(sheet.ncols, 100))))
                if sum(map(len, parts)) > TEXT_LIMIT: break
        wb.release_resources(); text='\n'.join(parts)
        warning = '读取缓存值，公式不重新计算。'
    else:
        office = shutil.which('libreoffice') or shutil.which('soffice')
        if not office: raise ValueError('此服务器暂未安装旧版 Office 转换器，请另存为 .docx 或 .pptx 后上传。')
        with tempfile.TemporaryDirectory() as temp:
            p=Path(temp); source=p/('input'+ext);source.write_bytes(data)
            target='docx' if ext=='.doc' else 'pptx'
            result=subprocess.run([office, '-env:UserInstallation='+ (p/'profile').as_uri(), '--headless', '--convert-to', target, '--outdir', temp, str(source)], capture_output=True, timeout=40)
            out=p/('input.'+target)
            if result.returncode or not out.exists(): raise ValueError('旧版 Office 转换失败，请另存为新格式。')
            text, warning=extract(out.name,out.read_bytes())
    if not text.strip(): raise ValueError('没有提取到文字。扫描件请先进行 OCR，或粘贴需要分析的文本。')
    if len(text)>TEXT_LIMIT: warning += ' 文本超过 180,000 字符，已截取；回答只依据选入上下文的部分。'
    return text[:TEXT_LIMIT], warning.strip()

